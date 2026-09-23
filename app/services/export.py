"""Turning a stored resume into a file its owner can take away (FR-5).

The export converts and nothing more: the "improved" resume is a version the user wrote
themselves from a match's suggestions, stored like any other (FR-2). Nothing here calls
a model, so nothing here can invent a fact the resume does not state.
"""

import io
import re
from typing import Literal

import docx

from app.models.resume import Resume

type ExportFormat = Literal["md", "docx"]

MEDIA_TYPES: dict[ExportFormat, str] = {
    "md": "text/markdown; charset=utf-8",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}

_XML_FORBIDDEN = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
"""Control characters XML 1.0 cannot carry; python-docx raises ValueError on them."""


def filename(resume: Resume, fmt: ExportFormat) -> str:
    """Name the file after the row, never after anything the user typed.

    The role or the original filename would put user text into Content-Disposition,
    where a quote or a newline breaks the header.
    """
    return f"resume-{resume.id}.{fmt}"


def to_markdown(resume: Resume) -> bytes:
    """Render the resume as Markdown: the role as a heading, then the text verbatim.

    Verbatim rather than escaped: it is the owner's own text, which may already use
    Markdown bullets, and escaping would turn those into literal asterisks.
    """
    heading = f"# {resume.target_role.strip()}\n\n" if resume.target_role else ""

    return f"{heading}{resume.content.strip()}\n".encode()


def to_docx(resume: Resume) -> bytes:
    """Render the resume as a Word document, one paragraph per line.

    The role is the heading, and blank lines stay blank paragraphs, so the spacing of
    the original survives. Control characters are dropped before the split, not after:
    splitlines() treats vertical tab, form feed and the separators U+001C-U+001E as
    line breaks, so each would otherwise leave a stray empty paragraph.
    """
    document = docx.Document()

    if resume.target_role:
        document.add_heading(_xml_safe(resume.target_role.strip()), level=1)

    for line in _xml_safe(resume.content).strip().splitlines():
        document.add_paragraph(line)

    buffer = io.BytesIO()
    document.save(buffer)

    return buffer.getvalue()


def _xml_safe(text: str) -> str:
    """Drop the control characters text extracted from a PDF sometimes carries."""
    return _XML_FORBIDDEN.sub("", text)


RENDERERS = {"md": to_markdown, "docx": to_docx}
